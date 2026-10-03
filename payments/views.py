from django.db.models import Q
from django.shortcuts import render
from .models import Customer,Payment
from django.shortcuts import get_object_or_404, redirect
from datetime import date
from django.utils import timezone
from django.shortcuts import get_object_or_404
from django.http import HttpResponse
from urllib.parse import quote
from django.shortcuts import redirect, get_object_or_404
from django.utils import timezone
from django.http import HttpResponse
def delete_customer(request, id):

    customer = get_object_or_404(Customer, id=id)

    customer.delete()

    return redirect('customers')
def edit_customer(request, id):

    customer = get_object_or_404(Customer, id=id)

    if request.method == "POST":

        customer.customer_code = request.POST['customer_code']
        customer.name = request.POST['name']
        customer.phone = request.POST['phone']

        customer.save()

        return redirect('customers')

    return render(
        request,
        'edit_customer.html',
        {
            'customer': customer
        }
    )
from django.http import HttpResponse

from django.db.models import Q

def customers(request):
    query = request.GET.get('q')

    customers = Customer.objects.all()

    if query:
        customers = customers.filter(
            Q(customer_code__icontains=query) |
            Q(name__icontains=query) |
            Q(phone__icontains=query)
        )

    return render(
        request,
        'customers.html',
        {
            'customers': customers
        }
    )
from urllib.parse import quote

def mark_paid(request, payment_id):

    payment = get_object_or_404(Payment, id=payment_id)

    if payment.status != 'Paid':
        payment.status = 'Paid'
        payment.paid_date = timezone.now().date()
        payment.paid_time = timezone.now()
        payment.save()

    phone = "91" + payment.customer.phone

    message = f"""
🏆 RK Lucky Draw

Dear {payment.customer.name},

Customer ID: {payment.customer.customer_code}

Your payment of ₹{payment.amount} for {payment.month}/{payment.year} has been received successfully.

Thank you for your payment.

- RK Lucky Draw
"""

    whatsapp_url = f"https://wa.me/{phone}?text={quote(message)}"

    return redirect(whatsapp_url)
def search_customer(request):
    query = request.GET.get('q')

    if query:
        customers = Customer.objects.filter(
            Q(name__icontains=query) |
            Q(phone__icontains=query) |
            Q(customer_code__icontains=query)
        )
    else:
        customers = []

    return render(
        request,
        'search.html',
        {'customers': customers}

    )
def payment_history(request):
    query = request.GET.get('q')

    customer = None
    payments = []

    if query:
        customer = Customer.objects.filter(
            Q(customer_code__icontains=query) |
            Q(name__icontains=query) |
            Q(phone__icontains=query)
        ).first()

        if customer:
            payments = Payment.objects.filter(
                customer=customer
            ).order_by('year', 'month')

    return render(
        request,
        'payment_history.html',
        {
            'customer': customer,
            'payments': payments
        }
    )
from datetime import date
from .models import Customer, Payment
from datetime import date
from .models import Customer, Payment


def generate_monthly_payments():

    today = date.today()

    for customer in Customer.objects.all():

        exists = Payment.objects.filter(
            customer=customer,
            month=today.month,
            year=today.year
        ).exists()

        if not exists:
            Payment.objects.create(
                customer=customer,
                month=today.month,
                year=today.year,
                status='Not Paid'
            )


def dashboard(request):

    generate_monthly_payments()

    today = date.today()

    current_month = today.month
    current_year = today.year

    total_customers = Customer.objects.count()

    monthly_payments = Payment.objects.filter(
        month=current_month,
        year=current_year
    )

    paid_count = monthly_payments.filter(
        status='Paid'
    ).count()

    unpaid_count = monthly_payments.filter(
        status='Not Paid'
    ).count()

    expected_collection = total_customers * 3000
    received_collection = paid_count * 3000
    pending_collection = expected_collection - received_collection

    return render(
        request,
        'dashboard.html',
        {
            'total_customers': total_customers,
            'paid_count': paid_count,
            'unpaid_count': unpaid_count,
            'expected_collection': expected_collection,
            'received_collection': received_collection,
            'pending_collection': pending_collection,
            'current_month': current_month,
            'current_year': current_year,
        }
    )
from django.db.models import Q
from datetime import date

def unpaid_customers(request):

    today = date.today()

    current_month = today.month
    current_year = today.year

    query = request.GET.get('q')

    payments = Payment.objects.filter(
        status='Not Paid',
        month=current_month,
        year=current_year
    )

    if query:
        payments = payments.filter(
            Q(customer__customer_code__icontains=query) |
            Q(customer__name__icontains=query) |
            Q(customer__phone__icontains=query)
        )
        print("SEARCH =", query)
        print("RESULTS =", payments.count())
        
        # for payment in payments:
        # payment.month_name = months[payment.month]
    return render(
        request,
        'unpaid_customers.html',
        {
            'payments': payments,
            'current_month': current_month,
            'current_year': current_year
        }
    )
from openpyxl import Workbook
from django.http import HttpResponse
from .models import Customer, Payment


def export_backup(request):

    wb = Workbook()

    # Customers Sheet
    ws1 = wb.active
    ws1.title = "Customers"

    ws1.append([
        "Customer Code",
        "Name",
        "Phone"
    ])

    for c in Customer.objects.all():

        ws1.append([
            c.customer_code,
            c.name,
            c.phone
        ])

    # Payments Sheet
    ws2 = wb.create_sheet("Payments")

    ws2.append([
    "Customer Code",
    "Customer Name",
    "Phone",
    "Month",
    "Year",
    "Status"
])

    for p in Payment.objects.all():

        ws2.append([
        p.customer.customer_code,
        p.customer.name,
        p.customer.phone,
        p.month,
        p.year,
        p.status
    ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    today = date.today()

    response['Content-Disposition'] = (f'attachment; filename=RK_Backup_{today}.xlsx')

    wb.save(response)

    return response
from urllib.parse import quote
from django.shortcuts import get_object_or_404, redirect

def send_reminder(request, payment_id):

    payment = get_object_or_404(Payment, id=payment_id)

    phone = "91" + payment.customer.phone

    months = {
        1:"January", 2:"February", 3:"March", 4:"April",
        5:"May", 6:"June", 7:"July", 8:"August",
        9:"September", 10:"October", 11:"November", 12:"December"
    }

    message = f"""
🏆 RK Lucky Draw

Dear {payment.customer.name},

This is a friendly reminder that your payment of ₹{payment.amount} for {months[payment.month]} {payment.year} is still pending.

Kindly make the payment at your earliest convenience.

Thank you.

- RK Lucky Draw
"""

    whatsapp_url = f"https://wa.me/91{payment.customer.phone}?text={quote(message)}"

    return redirect(whatsapp_url)
from openpyxl import Workbook
from django.http import HttpResponse
from datetime import date
from .models import Payment

def export_monthly_report(request):

    today = date.today()

    wb = Workbook()

    ws = wb.active
    ws.title = "Monthly Report"

    ws.append([
        "Customer Code",
        "Customer Name",
        "Phone",
        "Month",
        "Year",
        "Amount",
        "Status"
    ])

    payments = Payment.objects.filter(
        month=today.month,
        year=today.year
    )

    for p in payments:

        ws.append([
            p.customer.customer_code,
            p.customer.name,
            p.customer.phone,
            p.month,
            p.year,
            p.amount,
            p.status
        ])

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

    response['Content-Disposition'] = (
        f'attachment; filename=Monthly_Report_{today.month}_{today.year}.xlsx'
    )

    wb.save(response)

    return response